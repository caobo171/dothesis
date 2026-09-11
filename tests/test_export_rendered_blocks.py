"""What the results renderer emits has to survive the trip into Word.

Two properties, both invisible until the renderer started producing blocks that
actually reach an export:

  - Its sentinels are internal markup. `<!--dt-rendered:begin …-->` tells the
    coherence and similarity checkers "this table is computed, not typed"; the
    docx writer had no comment branch, so they printed as body text and the one
    sharing a line with the caption swallowed the caption with it.
  - A table screenshot is captioned ABOVE it and a figure BELOW it. Every
    Vietnamese thesis template says so and it is a supervisor's first-page check.
"""
import base64
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "engine"))
pytest.importorskip("docx")

from docx import Document  # noqa: E402

from engine.utils.export_professional import export_docx_basic  # noqa: E402

_MD = """# Chương 4

<!--dt-rendered:begin kind=measurement_model sha=e3e425c7f9aa-->
**Bảng 4.1 — Mô hình đo lường**

| Khái niệm | CR |
|---|---|
| ATT | 0.911 |
<!--dt-rendered:end kind=measurement_model-->
"""

# Smallest valid 1x1 PNG — python-docx has to be able to measure it.
_PNG = base64.b64decode(
    b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _png(path: Path) -> Path:
    path.write_bytes(_PNG)
    return path


def _picture_index(doc) -> int:
    return next(i for i, p in enumerate(doc.paragraphs)
                if p.runs and any("graphic" in r._r.xml for r in p.runs))


def test_sentinels_do_not_reach_the_docx(tmp_path):
    md, out = tmp_path / "c4.md", tmp_path / "c4.docx"
    md.write_text(_MD, encoding="utf-8")
    assert export_docx_basic(md, out)

    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "dt-rendered" not in body
    assert "<!--" not in body
    # The caption survives as its own paragraph, and the table still renders.
    assert any("Bảng 4.1" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1


def test_table_image_caption_sits_above_the_picture(tmp_path):
    img = _png(tmp_path / "hinh-04.png")
    md, out = tmp_path / "c4.md", tmp_path / "c4.docx"
    md.write_text(f"# Chương 4\n\n![Bảng 4.1 — Mô hình đo lường]({img})\n", encoding="utf-8")
    assert export_docx_basic(md, out)

    doc = Document(str(out))
    caption = next(i for i, p in enumerate(doc.paragraphs) if "Bảng 4.1" in p.text)
    assert caption < _picture_index(doc), "a table caption goes above the table"


def test_figure_image_caption_stays_below(tmp_path):
    img = _png(tmp_path / "model.png")
    md, out = tmp_path / "c3.md", tmp_path / "c3.docx"
    md.write_text(f"# Chương 3\n\n![Hình 3.1. Mô hình nghiên cứu]({img})\n", encoding="utf-8")
    assert export_docx_basic(md, out)

    doc = Document(str(out))
    caption = next(i for i, p in enumerate(doc.paragraphs) if "Hình 3.1" in p.text)
    assert _picture_index(doc) < caption, "a figure caption goes below the figure"
